"""Bounded local Excel ingestion. Rules: DQ-FILE-001, DQ-SCHEMA-001, DQ-VAL-001.

No network or database access. Error messages deliberately omit original cell values.
"""
from __future__ import annotations

import hashlib
import io
import zipfile
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import PurePath
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from lrm_data.dq import Finding
from lrm_data.schema import DATASET_MODELS, DatasetType, SourceMetadata, StrictRecord, expected_columns

MAX_FILE_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 160 * 1024 * 1024
MAX_COMPRESSION_RATIO = 120
MAX_DATA_ROWS = 100_000
MAX_META_ROWS = 50
META_FIELDS = tuple(SourceMetadata.model_fields)


@dataclass(frozen=True)
class ParsedRow:
    excel_row: int
    record: StrictRecord


@dataclass(frozen=True)
class LoadedSource:
    file_name: str
    file_hash: str
    metadata: SourceMetadata
    rows: tuple[ParsedRow, ...]


@dataclass(frozen=True)
class LoadResult:
    source: LoadedSource | None
    findings: tuple[Finding, ...]
    file_hash: str


def _failure(rule_id: str, dataset: str | None, row: int | None, field: str | None,
             message: str, file_hash: str) -> Finding:
    return Finding(rule_id, dataset, row, field, message, file_hash)


def _safe_filename(filename: str) -> bool:
    return (filename.endswith(".xlsx") and filename == PurePath(filename).name
            and not filename.startswith(".") and "\\" not in filename
            and "\n" not in filename and "\r" not in filename)


def _check_archive(blob: bytes) -> None:
    if len(blob) > MAX_FILE_BYTES:
        raise ValueError("File exceeds approved size limit")
    try:
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            infos = archive.infolist()
            if not infos or not any(i.filename == "xl/workbook.xml" for i in infos):
                raise ValueError("Not a valid XLSX workbook")
            total = 0
            for item in infos:
                total += item.file_size
                if total > MAX_UNCOMPRESSED_BYTES:
                    raise ValueError("Uncompressed XLSX size limit exceeded")
                if item.file_size > 1_000_000 and (
                    item.compress_size == 0 or
                    item.file_size > MAX_COMPRESSION_RATIO * item.compress_size
                ):
                    raise ValueError("Suspicious XLSX compression ratio")
                if (item.filename.startswith("/") or ".." in PurePath(item.filename).parts):
                    raise ValueError("Unsafe archive entry path")
                if "vbaproject" in item.filename.lower() or "externallinks" in item.filename.lower() or "connections.xml" in item.filename.lower():
                    raise ValueError("Macros, external links or connections are forbidden")
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise ValueError("Invalid XLSX archive") from exc


def _check_no_formulas_and_hidden(ws: Any) -> None:
    if ws.sheet_state != "visible":
        raise ValueError("Hidden sheets are forbidden")
    for dim in ws.row_dimensions.values():
        if dim.hidden:
            raise ValueError("Hidden worksheet rows are forbidden")
    for dim in ws.column_dimensions.values():
        if dim.hidden:
            raise ValueError("Hidden worksheet columns are forbidden")
    for row in ws:
        for cell in row:
            if cell.data_type == "f":
                raise ValueError("Formula cells are forbidden")
            if cell.data_type == "e":
                raise ValueError("Excel error cells are forbidden")


def _metadata(ws: Any) -> SourceMetadata:
    if ws.max_row > MAX_META_ROWS or ws.max_column != 2:
        raise ValueError("Invalid __meta__ layout")
    if [ws.cell(1, i).value for i in (1, 2)] != ["key", "value"]:
        raise ValueError("__meta__ must use key/value headers")
    values: dict[str, Any] = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        key, value = row
        if not isinstance(key, str) or key not in META_FIELDS or key in values:
            raise ValueError("Invalid/duplicate metadata key")
        if key == "reporting_date" and isinstance(value, datetime):
            if value.time().isoformat() != "00:00:00":
                raise ValueError("Reporting date contains non-midnight time")
            value = value.date()
        values[key] = value
    if set(values) != set(META_FIELDS):
        raise ValueError("Missing required source metadata")
    return SourceMetadata.model_validate(values)


def _headers(ws: Any, dataset: DatasetType) -> tuple[str, ...]:
    if ws.max_row < 2 or ws.max_column == 0:
        raise ValueError("No source records in data sheet")
    values = tuple(cell.value for cell in ws[1])
    if any(not isinstance(c, str) or not c for c in values):
        raise ValueError("Missing or invalid data column")
    if len(set(values)) != len(values):
        raise ValueError("Duplicate data column")
    if set(values) != set(expected_columns(dataset)):
        raise ValueError("Header does not match data contract schema 1.0")
    return values


def load_xlsx(file_name: str, payload: bytes) -> LoadResult:
    """Load known schema or return controlled DQ failures; input bytes never persisted."""
    sha = hashlib.sha256(payload).hexdigest()
    fail = lambda rid, typ, row, col, msg: _failure(rid, typ, row, col, msg, sha)
    if not _safe_filename(file_name):
        return LoadResult(None, (fail("DQ-FILE-001", None, None, None,
                                      "Only filename-safe .xlsx files are accepted"),), sha)
    try:
        _check_archive(payload)
        wb = load_workbook(io.BytesIO(payload), read_only=False, data_only=False,
                           keep_links=False)
        try:
            if set(wb.sheetnames) != {"__meta__", "data"}:
                raise ValueError("Expected exactly __meta__ and data worksheets")
            for sheet in wb:
                if sheet.title == "data" and sheet.max_row > MAX_DATA_ROWS + 1:
                    raise ValueError("Too many records for Phase 1 workbook")
                _check_no_formulas_and_hidden(sheet)
            meta = _metadata(wb["__meta__"])
            dataset = meta.dataset_type
            headers = _headers(wb["data"], dataset)
            model = DATASET_MODELS[dataset]
            findings: list[Finding] = []
            records: list[ParsedRow] = []
            for row_number, values in enumerate(
                wb["data"].iter_rows(min_row=2, values_only=True), start=2
            ):
                if all(v is None for v in values):
                    findings.append(fail("DQ-SCHEMA-001", dataset.value, row_number,
                                         None, "Blank source row in used range"))
                    continue
                row_data = dict(zip(headers, values, strict=True))
                try:
                    record = model.model_validate(row_data)
                    records.append(ParsedRow(row_number, record))
                except ValidationError as exc:
                    for error in exc.errors(include_url=False):
                        field = str(error["loc"][0]) if error["loc"] else None
                        rid = "DQ-AKAD-001" if field == "akad" else "DQ-VAL-001"
                        findings.append(fail(rid, dataset.value, row_number, field,
                                             "Invalid field value or type; consult schema 1.0"))
            return LoadResult(LoadedSource(file_name, sha, meta, tuple(records)),
                              tuple(findings), sha)
        finally:
            wb.close()
    except (ValueError, ValidationError, InvalidFileException, OSError, KeyError, TypeError) as exc:
        # Never propagate raw parser error text: it can include source cell values.
        rid = "DQ-FILE-001" if isinstance(exc, (InvalidFileException, OSError)) else "DQ-SCHEMA-001"
        return LoadResult(None, (fail(rid, None, None, None,
                                      "Workbook could not be safely parsed against schema 1.0"),), sha)
