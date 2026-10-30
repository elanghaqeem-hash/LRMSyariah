"""Generate demonstrably synthetic XLSX fixtures for 13 source contracts (WP1).

This script never reads bank files. All account/facility IDs carry a MASK-/TOKEN-
prefix. Synthetic reference ratios are NOT approved regulatory parameters.
"""
from __future__ import annotations

import argparse
import json
from datetime import date
from pathlib import Path
from typing import Any

from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lrm_data.schema import DATASET_MODELS, DatasetType, expected_columns

CUTOFF = date(2026, 10, 30)
META = {
    "schema_version": "1.0",
    "reporting_date": CUTOFF,
    "reporting_currency": "IDR",
    "source_system": "SYNTHETIC_GENERATOR",
    "extracted_at": "2026-10-30T01:00:00+07:00",
    "source_reference": "SYNTHETIC-WP1-NOT-BANK-DATA",
}

SAMPLES: dict[DatasetType, list[dict[str, Any]]] = {
    DatasetType.NERACA_HARIAN: [
        {"gl_code": f"GL-{i:03}", "pos_neraca": f"SYNTHETIC-GL-{i:03}",
         "saldo": str(i * 100000)} for i in range(1, 11)
    ],
    DatasetType.PENEMPATAN_BI: [{"placement_id":"BI-SYN-001", "jenis":"GWM",
        "gl_code":"GL-001", "akad":"wakalah", "saldo":"100000", "jatuh_tempo":None}],
    DatasetType.SURAT_BERHARGA_SYARIAH: [{"instrument_id":"SUK-SYN-001",
        "gl_code":"GL-002", "isin":"XS0000000000", "jenis":"SBSN",
        "issuer_code":"ISSUER-SYN-001", "akad":"ijarah", "rating":None,
        "agency":None, "jatuh_tempo":date(2027,12,30), "nilai_pasar":"200000",
        "encumbered_flag":False}],
    DatasetType.PENEMPATAN_BANK_LAIN: [{"placement_id":"PB-SYN-001",
        "gl_code":"GL-003", "counterparty_code":"BANK-SYN-001",
        "akad":"mudharabah", "jatuh_tempo":date(2026,11,20), "saldo":"300000"}],
    DatasetType.GIRO_TABUNGAN: [{"account_id":"MASK-GIRO0001", "gl_code":"GL-004",
        "product_code":"GIRO-SYN", "akad":"wadiah", "segmen":"perorangan",
        "lps_eligible":True, "operational_flag":False, "relationship_months":12,
        "saldo":"400000"}],
    DatasetType.DEPOSITO: [{"account_id":"MASK-DEPO0001", "gl_code":"GL-005",
        "product_code":"DEPO-SYN", "akad":"mudharabah", "jatuh_tempo":date(2027,1,30),
        "aro_flag":False, "early_withdrawal_penalty":"0", "segmen":"perorangan",
        "saldo":"500000"}],
    DatasetType.KEWAJIBAN_ANTARBANK: [{"funding_id":"KAB-SYN-001", "gl_code":"GL-006",
        "counterparty_code":"BANK-SYN-002", "akad":"wakalah",
        "jatuh_tempo":date(2026,11,18), "saldo":"600000"}],
    DatasetType.PEMBIAYAAN: [{"facility_id":"MASK-FAC00001", "gl_code":"GL-007",
        "product_code":"FIN-SYN", "akad":"murabahah", "kolektibilitas":1,
        "jatuh_tempo":date(2027,10,30), "jadwal_angsuran":json.dumps([
            {"tanggal":"2026-11-02", "nilai":"50000"},
            {"tanggal":"2026-12-02", "nilai":"50000"}]),
        "outstanding":"700000", "encumbered_flag":False}],
    DatasetType.KONTINJENSI: [{"commitment_id":"CONT-SYN-001", "gl_code":None,
        "jenis":"kafalah", "counterparty_type":"korporasi", "akad":"kafalah",
        "nilai":"800000"}],
    DatasetType.DISTRIBUSI_BAGI_HASIL: [{"distribution_id":"DIST-SYN-001",
        "product_code":"DEPO-SYN", "akad":"mudharabah", "tenor_months":3,
        "nisbah_nasabah":"0.60", "distributable_revenue":"90000",
        "equivalent_rate":"0.035"}],
    DatasetType.BENCHMARK_IMBAL_HASIL: [{"benchmark_type":"BENCH-SYN-001",
        "rate":"0.042", "sumber":"SYNTHETIC-NOT-MARKET-DATA",
        "tanggal_efektif":CUTOFF}],
    DatasetType.PERMODALAN: [{"component_id":"CAP-SYN-001", "gl_code":"GL-010",
        "komponen_modal":"MODAL-SYN", "nominal":"1000000"}],
    DatasetType.KEWAJIBAN_GWM_PLM: [{"reference_id":"GWM-TEST-UNVERIFIED",
        "rasio_gwm":"0", "rasio_plm":"0", "tanggal_efektif":CUTOFF,
        "sumber":"SYNTHETIC-NOT-REGULATORY", "legal_basis":"VERIFY-OJK-BI",
        "verification_status":"DRAFT_UNVERIFIED"}],
}


# Synthetic source transactions are all explicitly denominated in IDR.
for _rows in SAMPLES.values():
    for _record in _rows:
        _record["currency"] = "IDR"


def workbook_bytes(dataset: DatasetType, rows: list[dict[str, Any]] | None = None,
                   meta_overrides: dict[str, Any] | None = None) -> bytes:
    import io
    from openpyxl import Workbook
    wb = Workbook()
    meta = wb.active
    assert meta is not None
    meta.title = "__meta__"
    meta.append(["key", "value"])
    metadata = dict(META, dataset_type=dataset.value)
    if meta_overrides:
        metadata.update(meta_overrides)
    for key, value in metadata.items():
        meta.append([key, value])
    data = wb.create_sheet("data")
    cols = expected_columns(dataset)
    data.append(cols)
    for sheet in (meta, data):
        sheet.freeze_panes = "A2"
        sheet.row_dimensions[1].height = 26
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(fill_type="solid", fgColor="17365D")
            cell.alignment = Alignment(vertical="center")
    meta.column_dimensions["A"].width = 27
    meta.column_dimensions["B"].width = 47
    for col, field in enumerate(cols, start=1):
        data.column_dimensions[get_column_letter(col)].width = (
            42 if field == "jadwal_angsuran" else min(32, max(18, len(field) + 3))
        )
    for record in rows if rows is not None else SAMPLES[dataset]:
        data.append([record.get(field) for field in cols])
    output = io.BytesIO()
    wb.save(output)
    wb.close()
    return output.getvalue()


def main() -> None:
    parser = argparse.ArgumentParser(description="Generate synthetic WP1 .xlsx data contracts")
    parser.add_argument("--out", type=Path, default=Path("examples/synthetic"))
    args = parser.parse_args()
    args.out.mkdir(parents=True, exist_ok=True)
    for dataset in DATASET_MODELS:
        path = args.out / f"{dataset.value}.xlsx"
        path.write_bytes(workbook_bytes(dataset))
        print(f"Generated SYNTHETIC {path}")


if __name__ == "__main__":
    main()
